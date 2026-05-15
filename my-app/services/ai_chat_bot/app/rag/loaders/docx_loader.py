import asyncio
from pathlib import Path

from docx import Document as DocxDocument

from app.core.logger import log_error, log_event
from app.rag.exceptions.ingestion_exceptions import CorruptedDocumentError, EmptyDocumentError
from app.rag.loaders.base_loader import BaseDocumentLoader
from app.rag.models.document_models import DocumentMetadata, DocumentSection, RawDocument


class DOCXLoader(BaseDocumentLoader):
    loader_type = "python-docx"
    document_type = "docx"
    supported_extensions = {".docx"}

    async def load(self, file_path: str | Path) -> RawDocument:
        await self.validate_file(file_path)
        path = Path(file_path)

        try:
            paragraphs = await asyncio.to_thread(self._extract_paragraphs, path)
        except Exception as exc:
            log_error("RAG DOCX load failed", filename=path.name, error_type=type(exc).__name__)
            raise CorruptedDocumentError("docx extraction failed", details={"filename": path.name}) from exc

        sections = [
            DocumentSection(
                index=index,
                content=paragraph,
                section_type="paragraph",
                metadata={"source": "docx_paragraph"},
            )
            for index, paragraph in enumerate(paragraphs)
            if paragraph.strip()
        ]

        if not sections:
            raise EmptyDocumentError("docx contains no extractable text", details={"filename": path.name})

        metadata = await self.extract_metadata(path)
        document = RawDocument(content="\n\n".join(section.content for section in sections), sections=sections, metadata=metadata)

        log_event(
            "rag_docx_loaded",
            document_id=document.document_id,
            filename=metadata.filename,
            paragraph_count=metadata.paragraph_count,
        )
        return document

    async def extract_metadata(self, file_path: str | Path) -> DocumentMetadata:
        path = Path(file_path)
        paragraph_count = await asyncio.to_thread(lambda: len([p for p in self._extract_paragraphs(path) if p.strip()]))
        return await self.metadata_service.extract(
            path,
            loader_type=self.loader_type,
            document_type=self.document_type,
            paragraph_count=paragraph_count,
        )

    def _extract_paragraphs(self, path: Path) -> list[str]:
        document = DocxDocument(str(path))
        paragraphs = [paragraph.text or "" for paragraph in document.paragraphs]

        # Include simple table text because enterprise docs often keep policies,
        # pricing, or FAQs in tables.
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    paragraphs.append(" | ".join(cells))

        return paragraphs
